"""Public provider/CLI contracts; fake executables need no office app or Poppler."""

import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

from pypdf import PdfWriter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from engines import Settings, ProviderError, RenderSpec, render_pdf
from pdf_metrics import PageSpec
from matrix import MatrixSpec, ReferenceSpec

FAKE = r"""
import os, sys, shutil, json
from pathlib import Path
name = Path(sys.argv[0]).name
args = sys.argv[1:]
if '--version' in args or '-v' in args:
    print(name + ' fake 1.0'); sys.exit(0)
from pypdf import PdfReader
from PIL import Image
if name.startswith('deck') or name.startswith('office'):
    if 'sleep' in name:
        import time; time.sleep(20)
    if 'child' in name:
        import subprocess, time
        child = subprocess.Popen([sys.executable,'-c','import time; time.sleep(20)'])
        Path(os.environ['FAKE_CHILD_PID']).write_text(str(child.pid))
        time.sleep(20)
    src = Path(args[1]) if name.startswith('deck') else Path(args[-1])
    dst = Path(args[2]) if name.startswith('deck') else Path(args[args.index('--outdir')+1])/(src.stem+'.pdf')
    if 'link' in name:
        dst.symlink_to(os.environ['FAKE_PDF'])
    else:
        shutil.copyfile(os.environ['FAKE_PDF'], dst)
    if 'change' in name: src.write_bytes(b'changed while rendering')
    if 'invalid' in name: dst.write_bytes(b'not a PDF')
    if 'fail' in name: print('deliberate partial failure', file=sys.stderr); sys.exit(7)
elif name == 'pdffonts':
    print('name type encoding emb sub uni object ID')
    print('--------------------------------------')
    print('ABCDEF+ProbeSans TrueType WinAnsi yes yes yes 1 0')
elif name == 'pdftotext':
    n = len(PdfReader(args[-2]).pages)
    Path(args[-1]).write_text('<doc>' + ''.join('<page><flow><block><line><word xMin="10" yMin="20">Probe</word></line></block></flow></page>' for _ in range(n)) + '</doc>')
elif name.startswith('pdftoppm'):
    import math
    pdf=PdfReader(args[-2])
    dpi=float(args[args.index('-r')+1]) if '-r' in args else 150
    for i,page in enumerate(pdf.pages):
        size=(math.ceil(float(page.mediabox.width)*dpi/72),math.ceil(float(page.mediabox.height)*dpi/72))
        if '-scale-to-x' in args:
            size=(int(args[args.index('-scale-to-x')+1]),int(args[args.index('-scale-to-y')+1]))
        if 'bad' in name: size=(8,8)
        suffix='.png' if '-singlefile' in args else '-'+str(i+1)+'.png'
        Image.new('L' if '-gray' in args else 'RGB', size, 255 if '-gray' in args else 'white').save(args[-1]+suffix)
else: sys.exit(9)
"""


class Engines(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name).resolve()
        self.bin = self.root / "bin"
        self.bin.mkdir()
        for name in [
            "deck",
            "deck-fail",
            "deck-sleep",
            "deck-change",
            "deck-invalid",
            "deck-link",
            "deck-child",
            "office",
            "pdftoppm",
            "pdftoppm-bad",
            "pdftotext",
            "pdffonts",
        ]:
            p = self.bin / name
            p.write_text(f"#!{sys.executable}\n" + FAKE)
            p.chmod(0o755)
        self.pdf = self.root / "input.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=960, height=540)
        with self.pdf.open("wb") as stream:
            writer.write(stream)
        self.source = self.root / "input.pptx"
        self.source.write_bytes(b"synthetic provider input")
        self.font = self.root / "font.ttf"
        self.font.write_bytes(b"synthetic font fingerprint")
        self.old = dict(os.environ)
        for key in list(os.environ):
            if key.startswith("PARITY_"):
                del os.environ[key]
        os.environ.update(
            FAKE_PDF=str(self.pdf), XDG_CONFIG_HOME=str(self.root / "xdg")
        )
        self.config = self.root / "config.toml"
        self.config.write_text(f'''[providers.deckcraft]
binary = "{self.bin / "deck"}"
revision = "test-revision"
[providers.libreoffice]
binary = "{self.bin / "office"}"
[tools]
pdftoppm = "{self.bin / "pdftoppm"}"
pdftotext = "{self.bin / "pdftotext"}"
pdffonts = "{self.bin / "pdffonts"}"
[fonts]
files = ["{self.font}"]
expected = ["ProbeSans"]
''')
        self.settings = Settings.load(self.config)

    def tearDown(self):
        os.environ.clear()
        os.environ.update(self.old)
        self.tmp.cleanup()

    def render(self, **kwargs):
        return render_pdf(
            "deckcraft",
            self.source,
            self.root / "cache",
            self.settings,
            RenderSpec(pages=PageSpec(count=1), **kwargs),
        )

    def failed_render(self, mode, settings, pages=1):
        cache = self.root / f"cache-{mode}"
        with self.assertRaises(ProviderError) as ctx:
            render_pdf(
                "deckcraft",
                self.source,
                cache,
                settings,
                RenderSpec(pages=PageSpec(count=pages)),
            )
        self.assertFalse(list(cache.glob("*/manifest.json")))
        return ctx.exception.code

    def test_concurrent_publication_reuses_a_verified_winner(self):
        from concurrent.futures import ThreadPoolExecutor

        with ThreadPoolExecutor(max_workers=2) as workers:
            results = list(workers.map(lambda _: self.render(), range(2)))
        self.assertEqual(results[0]["key"], results[1]["key"])
        self.assertEqual(sum(r["cache_hit"] for r in results), 1)
        self.assertTrue(self.render()["cache_hit"])

    def test_render_and_hash_validated_cache(self):
        first = self.render()
        second = self.render()
        self.assertFalse(first["cache_hit"])
        self.assertTrue(second["cache_hit"])
        self.assertEqual(first["key"], second["key"])
        self.assertEqual(first["fonts"][0]["name"], "ProbeSans")
        Path(first["pdf"]).write_bytes(b"corrupt")
        with self.assertRaisesRegex(ProviderError, "cache"):
            self.render()

    def test_corrupt_manifest_cannot_claim_a_cache_hit(self):
        result = self.render()
        manifest = Path(result["directory"]) / "manifest.json"
        original = manifest.read_bytes()
        for bad in [
            "not JSON",
            json.dumps({"key": result["key"]}),
            json.dumps({**json.loads(original), "artifacts": {}}),
        ]:
            manifest.write_text(bad)
            with self.assertRaises(ProviderError) as ctx:
                self.render()
            self.assertEqual(ctx.exception.code, "cache_corrupt")
        manifest.write_bytes(original)
        self.assertTrue(self.render()["cache_hit"])

    def test_source_font_revision_and_binary_changes_invalidate_cache(self):
        keys = [self.render()["key"]]
        self.source.write_bytes(b"changed source")
        keys.append(self.render()["key"])
        self.font.write_bytes(b"changed font")
        keys.append(self.render()["key"])
        self.settings.data["providers"]["deckcraft"]["revision"] = "another-revision"
        keys.append(self.render()["key"])
        with (self.bin / "deck").open("a") as f:
            f.write("\n# new binary revision\n")
        keys.append(self.render()["key"])
        renamed = self.source.with_name("renamed.pptx")
        renamed.write_bytes(self.source.read_bytes())
        keys.append(
            render_pdf(
                "deckcraft",
                renamed,
                self.root / "cache",
                self.settings,
                RenderSpec(pages=PageSpec(count=1)),
            )["key"]
        )
        self.assertEqual(len(set(keys)), len(keys))

    def test_explicit_config_precedes_env_and_cli_precedes_config(self):
        os.environ["PARITY_DECKCRAFT_BIN"] = "/missing-env"
        self.render()
        with self.assertRaisesRegex(ProviderError, "executable"):
            self.render(override="/missing-cli")
        self.settings.data["providers"]["deckcraft"]["sha256"] = "0" * 64
        with self.assertRaisesRegex(ProviderError, "hash"):
            self.render()

    def test_rasterizer_cannot_silently_change_recorded_dimensions(self):
        self.settings.data["tools"]["pdftoppm"] = str(self.bin / "pdftoppm-bad")
        self.assertEqual(
            self.failed_render("raster-dimensions", self.settings), "geometry"
        )

    def test_pptx_cli_discovery_prefers_new_name_without_overriding_pins(self):
        renamed = self.bin / "pptx-cli"
        renamed.write_text(FAKE)
        renamed.chmod(0o755)
        self.settings.data["providers"]["deckcraft"]["binary"] = ""
        with mock.patch.dict(os.environ, {"PATH": str(self.bin)}):
            self.assertEqual(self.settings.binary("deckcraft"), renamed)
            self.settings.data["providers"]["deckcraft"]["binary"] = str(
                self.bin / "deck"
            )
            self.assertEqual(self.settings.binary("deckcraft"), self.bin / "deck")
            self.assertEqual(self.settings.binary("deckcraft", str(renamed)), renamed)

    def test_selected_missing_tool_does_not_fall_back(self):
        self.settings.data["tools"]["pdffonts"] = "/missing-selected-tool"
        with self.assertRaisesRegex(ProviderError, "executable"):
            self.render()

    def test_failure_timeout_font_mismatch_and_page_count_publish_no_cache(self):
        for mode, code in [
            ("fail", "provider_failed"),
            ("sleep", "timeout"),
            ("font", "font_mismatch"),
            ("pages", "page_count"),
        ]:
            with self.subTest(mode=mode):
                settings = Settings.load(self.config)
                settings.data["timeout_seconds"] = 0.5 if mode == "sleep" else 10
                if mode in ("fail", "sleep"):
                    settings.data["providers"]["deckcraft"]["binary"] = str(
                        self.bin / f"deck-{mode}"
                    )
                if mode == "font":
                    settings.data["fonts"]["expected"] = ["MissingFace"]
                self.assertEqual(
                    self.failed_render(mode, settings, 2 if mode == "pages" else 1),
                    code,
                )

    def test_missing_config_is_an_error_not_an_auto_created_override(self):
        with self.assertRaises(ProviderError):
            Settings.load(self.root / "missing.toml")
        self.assertFalse((self.root / "missing.toml").exists())

    def test_matrix_supplied_reference_and_cli_are_attributed(self):
        from matrix import execute

        out = self.root / "evidence"
        report = execute(
            self.source,
            out,
            self.root / "cache",
            self.settings,
            MatrixSpec(
                engines=("deckcraft", "libreoffice"),
                reference=self.pdf,
                reference_engine="powerpoint",
                reference_version="provided-test",
                pages=PageSpec(count=1),
            ),
        )
        self.assertEqual(report["status"], "complete")
        self.assertEqual(len(report["pairs"]), 3)
        self.assertEqual(report["providers"]["powerpoint"]["origin"], "supplied_pdf")
        self.assertEqual(
            self.pdf.read_bytes(), (out / "powerpoint" / "source.pdf").read_bytes()
        )
        self.assertTrue((out / "report.json").exists())
        with self.assertRaisesRegex(ProviderError, "exists"):
            execute(
                self.source,
                out,
                self.root / "cache",
                self.settings,
                MatrixSpec(engines=("deckcraft",), pages=PageSpec(count=1)),
            )

    def test_cli_json_failure_and_legacy_help(self):
        cli = Path(__file__).resolve().parents[1] / "parity.py"
        p = subprocess.run(
            [
                sys.executable,
                str(cli),
                "matrix",
                "--pptx",
                str(self.source),
                "--output",
                str(self.root / "out"),
                "--config",
                str(self.config),
                "--deckcraft-bin",
                "/absent",
                "--engines",
                "deckcraft",
            ],
            capture_output=True,
            text=True,
        )
        self.assertNotEqual(p.returncode, 0)
        self.assertIn(json.loads(p.stdout)["status"], ["error", "partial"])
        help_result = subprocess.run(
            [sys.executable, str(cli), "ingest", "--help"],
            capture_output=True,
            text=True,
        )
        self.assertEqual(help_result.returncode, 0)
        self.assertIn("--slides", help_result.stdout)

    def test_immutable_reference_and_legacy_snapshot(self):
        from matrix import ingest_reference

        deck_dir = self.root / "references"
        reference = ingest_reference(
            self.source,
            deck_dir,
            self.root / "cache",
            ReferenceSpec(
                reference=self.pdf,
                reference_engine="powerpoint",
                config=self.config,
                pages=PageSpec(count=1),
            ),
        )
        original = reference.read_bytes()
        meta = json.loads((deck_dir / "reference.json").read_text())
        self.assertEqual(meta["renderer"], "powerpoint")
        with self.assertRaisesRegex(ProviderError, "Reference"):
            ingest_reference(
                self.source,
                deck_dir,
                self.root / "cache",
                ReferenceSpec(
                    renderer="deckcraft", config=self.config, pages=PageSpec(count=1)
                ),
            )
        self.assertEqual(reference.read_bytes(), original)
        legacy = self.root / "legacy"
        legacy.mkdir()
        (legacy / "original.pdf").write_bytes(original)
        self.assertEqual(
            ingest_reference(self.source, legacy, self.root / "cache"),
            legacy / "original.pdf",
        )
        with self.assertRaises(ProviderError):
            ingest_reference(
                self.source,
                legacy,
                self.root / "cache",
                ReferenceSpec(renderer="deckcraft"),
            )

    def test_one_provider_failure_still_publishes_attributed_partial_evidence(self):
        from matrix import execute

        self.settings.data["providers"]["deckcraft"]["binary"] = str(
            self.bin / "deck-fail"
        )
        out = self.root / "mixed"
        report = execute(
            self.source,
            out,
            self.root / "cache",
            self.settings,
            MatrixSpec(reference=self.pdf, pages=PageSpec(count=1)),
        )
        self.assertEqual(report["status"], "error")
        self.assertEqual(
            report["providers"]["deckcraft"]["error"]["code"], "provider_failed"
        )
        self.assertEqual(report["providers"]["libreoffice"]["status"], "ok")
        self.assertEqual(len(report["pairs"]), 1)
        self.assertTrue((out / "report.json").exists())

    def test_missing_optional_engine_is_partial_not_success(self):
        from matrix import execute

        self.settings.data["providers"]["libreoffice"]["binary"] = ""
        self.settings.explicit["providers"]["libreoffice"].pop("binary")
        os.environ["PATH"] = str(self.bin)
        report = execute(
            self.source,
            self.root / "partial",
            self.root / "cache",
            self.settings,
            MatrixSpec(
                engines=("deckcraft", "libreoffice"),
                reference=self.pdf,
                pages=PageSpec(count=1),
            ),
        )
        self.assertEqual(report["status"], "partial")
        self.assertEqual(report["providers"]["libreoffice"]["status"], "skipped")

    def test_matrix_freezes_common_input_and_never_passes_the_original_to_provider(
        self,
    ):
        from matrix import execute

        original = self.source.read_bytes()
        self.settings.data["providers"]["deckcraft"]["binary"] = str(
            self.bin / "deck-change"
        )
        with self.assertRaises(ProviderError) as ctx:
            execute(
                self.source,
                self.root / "changed-run",
                self.root / "cache",
                self.settings,
                MatrixSpec(engines=("deckcraft",), pages=PageSpec(count=1)),
            )
        self.assertEqual(ctx.exception.code, "inputs_changed")
        self.assertEqual(self.source.read_bytes(), original)
        self.assertFalse((self.root / "changed-run").exists())

    def test_no_font_pack_disables_cache_reuse(self):
        self.settings.data["fonts"]["files"] = []
        self.render()
        result = self.render()
        self.assertFalse(result["cache_hit"])
        self.assertFalse(result["cacheable"])
        self.settings.data["fonts"]["files"] = [str(self.font)]
        self.settings.data["cache_enabled"] = False
        self.render()
        self.assertFalse(self.render()["cache_hit"])

    def test_bad_outputs_and_inputs_changed_during_run_do_not_publish(self):
        for mode, code in [
            ("invalid", "invalid_pdf"),
            ("link", "provider_output"),
            ("change", "inputs_changed"),
        ]:
            with self.subTest(mode=mode):
                settings = Settings.load(self.config)
                settings.data["providers"]["deckcraft"]["binary"] = str(
                    self.bin / f"deck-{mode}"
                )
                self.assertEqual(self.failed_render(mode, settings), code)

    def test_aspect_ratio_mismatch_and_git_publication_guard(self):
        from matrix import execute

        writer = PdfWriter()
        writer.add_blank_page(width=720, height=540)
        with self.pdf.open("wb") as stream:
            writer.write(stream)
        with self.assertRaises(ProviderError) as ctx:
            self.render()
        self.assertEqual(ctx.exception.code, "aspect_ratio")
        destination = self.root / "tracked"
        (destination / ".git").mkdir(parents=True)
        with self.assertRaisesRegex(ProviderError, "outside git"):
            execute(
                self.source,
                destination / "run",
                self.root / "cache",
                self.settings,
                MatrixSpec(engines=("deckcraft",), pages=PageSpec(count=1)),
            )

    @unittest.skipUnless(os.name == "posix", "POSIX process-group supervision")
    def test_cancellation_reaps_the_isolated_command(self):
        from unittest.mock import Mock, patch
        from engines import process

        child = Mock(pid=987654321)
        child.wait.side_effect = [KeyboardInterrupt(), None]
        with (
            patch("engines.subprocess.Popen", return_value=child),
            patch("engines.os.killpg") as kill,
        ):
            with self.assertRaises(KeyboardInterrupt):
                process(["fake"], 10)
            self.assertEqual(child.wait.call_count, 2)
            kill.assert_called_once()

    @unittest.skipUnless(os.name == "posix", "POSIX process-group supervision")
    def test_timeout_kills_descendant_not_the_host_session(self):
        pid_file = self.root / "child.pid"
        os.environ["FAKE_CHILD_PID"] = str(pid_file)
        self.settings.data["providers"]["deckcraft"]["binary"] = str(
            self.bin / "deck-child"
        )
        self.settings.data["timeout_seconds"] = 2
        with self.assertRaises(ProviderError) as ctx:
            self.render()
        self.assertEqual(ctx.exception.code, "timeout")
        child = int(pid_file.read_text())
        state = subprocess.run(
            ["ps", "-p", str(child), "-o", "stat="], capture_output=True, text=True
        )
        self.assertTrue(
            state.returncode != 0 or state.stdout.strip().startswith("Z"), state.stdout
        )

    def test_complete_cli_and_ingest_use_the_same_provider_contract(self):
        from pptx import Presentation

        p = Presentation()
        p.slide_width = 12192000
        p.slide_height = 6858000
        p.slides.add_slide(p.slide_layouts[6])
        p.save(self.source)
        cli = Path(__file__).resolve().parents[1] / "parity.py"
        env = dict(
            os.environ,
            PARITY_LAB=str(self.root / "lab"),
            PATH=str(self.bin) + os.pathsep + os.environ["PATH"],
        )
        matrix = subprocess.run(
            [
                sys.executable,
                str(cli),
                "matrix",
                "--pptx",
                str(self.source),
                "--output",
                str(self.root / "cli"),
                "--cache",
                str(self.root / "cache"),
                "--config",
                str(self.config),
            ],
            env=env,
            capture_output=True,
            text=True,
        )
        self.assertEqual(matrix.returncode, 0, matrix.stdout + matrix.stderr)
        self.assertEqual(json.loads(matrix.stdout)["status"], "complete")
        ingest = subprocess.run(
            [
                sys.executable,
                str(cli),
                "ingest",
                str(self.source),
                "--renderer",
                "deckcraft",
                "--config",
                str(self.config),
                "--case-prefix",
                "test",
            ],
            env=env,
            capture_output=True,
            text=True,
        )
        self.assertEqual(ingest.returncode, 0, ingest.stdout + ingest.stderr)
        case = json.loads((self.root / "lab/cases/test-s01/case.json").read_text())
        self.assertEqual(case["renderer"], "deckcraft")
        self.assertIn("provenance", case["renderer_provenance"])

    def test_global_config_creation_env_priority_and_unknown_key_rejection(self):
        previous = Path.cwd()
        try:
            os.chdir(self.root)
            settings = Settings.load()
            self.assertTrue((self.root / "xdg/sldr-parity/config.toml").exists())
            os.environ["PARITY_DECKCRAFT_BIN"] = str(self.bin / "deck")
            selected = settings.binary("deckcraft")
            self.assertEqual(selected, self.bin / "deck")
            self.config.write_text("unknown_key = 1")
            with self.assertRaisesRegex(ProviderError, "Unknown config key"):
                Settings.load(self.config)
        finally:
            os.chdir(previous)
