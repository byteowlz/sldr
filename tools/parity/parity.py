#!/usr/bin/env -S uv run --quiet --script
# /// script
# requires-python = ">=3.11"
# dependencies = ["python-pptx>=1.0", "scikit-image>=0.24", "pillow>=10", "numpy>=1.26"]
# ///
"""parity — measure how closely sldr can recreate a real PowerPoint slide.

Deterministic tools for the recreation loop; the agent (pi with the
sldr-parity skill) does the recreating. One folder per case:

    <lab>/cases/<case>/
      case.json      source deck, slide number, sldr target (library + slide)
      original.png   the slide as rendered from the source deck (1920x1080)
      extract.json   shapes, geometry (% of slide), text runs, fills, theme
      media/         pictures from the slide, as embedded in the deck
      lib/           the case's own sldr library (slides, playlists, flavors, layouts)
      config/        XDG config pointing sldr at lib/ (never the user's library)
      out/           sldr output
      render.png  compare.png  heatmap.png  score.json  history.jsonl  gaps.jsonl

Commands
  ingest  DECK --slides 1,4-6 [--case-prefix NAME]   create cases from a deck
  baseline CASE --library DIR --slide NAME [--flavor F]  score an existing hand port
  render  CASE        build the case's sldr slide and rasterize it
  score   CASE        render, then compare with the original (SSIM, heatmap, text recall)
  gap     CASE --kind K --feature F --evidence E    record one remaining difference
  report              HTML overview of every case in the lab

The lab defaults to ~/sldr-lab (PARITY_LAB overrides). It holds source decks
and their renders: keep it private, never in a repository.
"""

from __future__ import annotations

import argparse
import datetime as dt
import html
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw, ImageFont

W, H = 1920, 1080
LAB = Path(os.environ.get("PARITY_LAB", "~/sldr-lab")).expanduser()
GAP_KINDS = ("missing_feature", "parameter", "agent_error", "renderer")


# ---------------------------------------------------------------- utilities


def die(msg: str) -> None:
    print(f"parity: {msg}", file=sys.stderr)
    sys.exit(1)


def case_dir(name: str) -> Path:
    p = Path(name)
    d = p if p.is_dir() else LAB / "cases" / name
    if not (d / "case.json").exists():
        die(f"no case at {d}")
    return d


def load(p: Path) -> dict:
    return json.loads(p.read_text())


def save(p: Path, data) -> None:
    p.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n")


def parse_slides(spec: str, total: int) -> list[int]:
    out: list[int] = []
    for part in spec.split(","):
        part = part.strip()
        if part == "all":
            return list(range(1, total + 1))
        if "-" in part:
            a, b = part.split("-")
            out.extend(range(int(a), int(b) + 1))
        elif part:
            out.append(int(part))
    bad = [n for n in out if not 1 <= n <= total]
    if bad:
        die(f"slides {bad} outside 1..{total}")
    return out


def run(cmd: list[str], env: dict | None = None, timeout: int = 600) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, env=env, capture_output=True, text=True, timeout=timeout)


def pdf_page_png(pdf: Path, page: int, out: Path) -> None:
    stem = out.with_suffix("")
    r = run(["pdftoppm", "-png", "-singlefile", "-f", str(page), "-l", str(page),
             "-scale-to-x", str(W), "-scale-to-y", str(H), str(pdf), str(stem)])
    if r.returncode != 0 or not out.exists():
        die(f"pdftoppm failed for page {page}: {r.stderr.strip()}")


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:40] or "deck"


# ---------------------------------------------------------------- extraction

EMU_PT = 12700


def color_hex(color) -> str | None:
    try:
        if color is not None and color.type is not None and color.rgb is not None:
            return "#" + str(color.rgb)
    except Exception:
        pass
    try:
        if color is not None and color.theme_color is not None:
            return f"theme:{color.theme_color}".lower().replace("msothemecolorindex.", "")
    except Exception:
        pass
    return None


def fill_of(shape) -> str | None:
    try:
        f = shape.fill
        if f.type == 1:  # MSO_FILL.SOLID
            return color_hex(f.fore_color)
        if f.type is not None:
            return str(f.type).split(".")[-1].split(" ")[0].lower()
    except Exception:
        pass
    return None


def text_of(tf) -> list[dict]:
    paras = []
    for p in tf.paragraphs:
        runs = []
        for r in p.runs:
            f = r.font
            runs.append({
                "text": r.text,
                "size_pt": f.size.pt if f.size else None,
                "bold": f.bold,
                "italic": f.italic,
                "font": f.name,
                "color": color_hex(f.color) if f.color and f.color.type is not None else None,
            })
        if runs or p.text.strip():
            paras.append({
                "level": p.level,
                "align": str(p.alignment).split(".")[-1].split(" ")[0].lower() if p.alignment else None,
                "runs": runs,
            })
    return paras


R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def save_pictures(shape, sid: str, media: Path) -> list[dict]:
    """Every picture a shape embeds, best first: an SVG (svgBlip) before its
    raster fallback. python-pptx's .image misses SVG pictures and some
    fallback forms, so follow the r:embed relationships directly."""
    out = []
    try:
        el = shape._element
        rids = [b.get(R_NS + "embed") for b in el.iter() if b.tag.endswith("}svgBlip")]
        rids += [b.get(R_NS + "embed") for b in el.iter() if b.tag.endswith("}blip")]
        seen = set()
        for i, rid in enumerate(r for r in rids if r):
            if rid in seen:
                continue
            seen.add(rid)
            part = shape.part.related_part(rid)
            ext = part.partname.ext
            name = f"{sid.replace('.', '_')}{'' if i == 0 else f'-{i}'}.{ext}"
            (media / name).write_bytes(part.blob)
            rec = {"file": f"media/{name}"}
            try:
                from PIL import Image as PILImage
                import io
                rec["px"] = list(PILImage.open(io.BytesIO(part.blob)).size)
            except Exception:
                pass
            out.append(rec)
    except Exception:
        pass
    return out


def walk(shapes, sw: int, sh: int, media: Path, prefix: str = "") -> list[dict]:
    out = []
    for i, s in enumerate(shapes):
        sid = f"{prefix}{i}"
        kind = str(s.shape_type).split(".")[-1].split(" ")[0].lower() if s.shape_type else "unknown"
        item: dict = {"id": sid, "name": s.name, "kind": kind}
        try:
            item["box"] = {k: round(v, 2) for k, v in {
                "x": s.left / sw * 100, "y": s.top / sh * 100,
                "w": s.width / sw * 100, "h": s.height / sh * 100}.items()}
        except Exception:
            item["box"] = None
        if getattr(s, "rotation", 0):
            item["rotation"] = s.rotation
        if s.is_placeholder:
            try:
                item["placeholder"] = str(s.placeholder_format.type).split(".")[-1].split(" ")[0].lower()
            except Exception:
                pass
        fill = fill_of(s) if hasattr(s, "fill") else None
        if fill:
            item["fill"] = fill
        if s.has_text_frame and s.text_frame.text.strip():
            item["text"] = text_of(s.text_frame)
        if kind == "picture" or hasattr(s, "image"):
            saved = save_pictures(s, sid, media)
            if saved:
                item["image"] = saved[0]
                if len(saved) > 1:
                    item["image_alternates"] = saved[1:]
                try:
                    c = s.crop_left, s.crop_top, s.crop_right, s.crop_bottom
                    if any(c):
                        item["image"]["crop"] = [round(v, 4) for v in c]
                except Exception:
                    pass
        if getattr(s, "has_table", False) and s.has_table:
            item["table"] = [[cell.text for cell in row.cells] for row in s.table.rows]
        if getattr(s, "has_chart", False) and s.has_chart:
            item["chart"] = str(s.chart.chart_type).split(".")[-1].split(" ")[0].lower()
        if kind == "group":
            item["children"] = walk(s.shapes, sw, sh, media, sid + ".")
        out.append(item)
    return out


def theme_of(pptx_path: Path) -> dict:
    """Theme colors and fonts straight from the first theme part."""
    theme: dict = {}
    try:
        with zipfile.ZipFile(pptx_path) as z:
            name = sorted(n for n in z.namelist() if re.match(r"ppt/theme/theme\d+\.xml$", n))[0]
            xml = z.read(name).decode("utf8", "replace")
        scheme = re.search(r"<a:clrScheme.*?</a:clrScheme>", xml, re.S)
        if scheme:
            for slot, val in re.findall(r"<a:(dk1|lt1|dk2|lt2|accent\d|hlink|folHlink)>.*?(?:srgbClr val|lastClr)=\"([0-9A-Fa-f]{6})\"", scheme.group(0), re.S):
                theme[slot] = "#" + val
        for role in ("majorFont", "minorFont"):
            m = re.search(rf"<a:{role}>\s*<a:latin typeface=\"([^\"]*)\"", xml)
            if m:
                theme[role] = m.group(1)
    except Exception:
        pass
    return theme


def background_image(part_owner, media: Path, name: str) -> str | None:
    """The picture behind a slide, layout or master (<p:bg> with a blip), saved to media/."""
    try:
        el = part_owner._element
        blips = el.xpath("./p:cSld/p:bg//a:blip")
        if not blips:
            return None
        rid = blips[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        image_part = part_owner.part.related_part(rid)
        ext = image_part.partname.ext
        (media / f"{name}.{ext}").write_bytes(image_part.blob)
        return f"media/{name}.{ext}"
    except Exception:
        return None


def background_color(part_owner) -> str | None:
    try:
        el = part_owner._element
        clr = el.xpath("./p:cSld/p:bg//a:solidFill/a:srgbClr")
        return "#" + clr[0].get("val") if clr else None
    except Exception:
        return None


# ---------------------------------------------------------------- ingest


def render_deck_pdf(deck: Path, deck_dir: Path) -> Path:
    """Render the whole deck once (LibreOffice) and cache it."""
    pdf = deck_dir / "original.pdf"
    if pdf.exists() and pdf.stat().st_mtime >= deck.stat().st_mtime:
        return pdf
    deck_dir.mkdir(parents=True, exist_ok=True)
    tmp = deck_dir / "lo"
    tmp.mkdir(exist_ok=True)
    r = run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(deck)], timeout=900)
    produced = tmp / (deck.stem + ".pdf")
    if r.returncode != 0 or not produced.exists():
        die(f"LibreOffice could not render {deck}: {r.stderr.strip()[:300]}")
    shutil.move(produced, pdf)
    shutil.rmtree(tmp, ignore_errors=True)
    return pdf


def write_case_config(case: Path, library: Path, layout_dir: Path) -> None:
    cfg = case / "config" / "sldr"
    cfg.mkdir(parents=True, exist_ok=True)
    seeds = Path("~/.config/sldr/flavors").expanduser()
    (cfg / "config.toml").write_text(
        "# Generated by parity: this case's sldr sees only its own library.\n"
        "[config]\n"
        f'library = "{library}"\n'
        f'flavor_dir = "{seeds}"\n'
        f'layout_dir = "{layout_dir}"\n'
        f'scaffold_dir = "{case / "config" / "scaffolds"}"\n'
        'default_flavor = "default"\n\n'
        "[presentations]\n"
        f'slide_dir = "{library / "slides"}"\n'
        f'playlist_dir = "{library / "playlists"}"\n'
        f'output_dir = "{case / "out"}"\n'
    )


def cmd_ingest(a) -> None:
    from pptx import Presentation

    deck = Path(a.deck).expanduser().resolve()
    prs = Presentation(str(deck))
    total = len(prs.slides)
    numbers = parse_slides(a.slides, total)
    prefix = a.case_prefix or slug(deck.stem)
    deck_dir = LAB / "decks" / prefix
    pdf = render_deck_pdf(deck, deck_dir)
    theme = theme_of(deck)
    sw, sh = prs.slide_width, prs.slide_height
    for n in numbers:
        name = f"{prefix}-s{n:02d}"
        case = LAB / "cases" / name
        if case.exists() and not a.force:
            print(f"  = {name} (exists; --force to recreate)")
            continue
        if case.exists():
            shutil.rmtree(case)
        (case / "media").mkdir(parents=True)
        lib = case / "lib"
        for sub in ("slides", "playlists", "flavors", "layouts"):
            (lib / sub).mkdir(parents=True)
        slide = prs.slides[n - 1]
        extract = {
            "deck": str(deck), "slide": n, "of": total,
            "size": {"w_emu": sw, "h_emu": sh, "aspect": round(sw / sh, 4)},
            "layout": slide.slide_layout.name, "master_layout": slide.slide_layout.slide_master.name if hasattr(slide.slide_layout.slide_master, "name") else None,
            "theme": theme,
            "notes": slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else None,
            "shapes": walk(slide.shapes, sw, sh, case / "media"),
            # Inherited from the slide layout and master: house-style art,
            # logos and footers usually live here, not on the slide.
            "layout_shapes": walk(slide.slide_layout.shapes, sw, sh, case / "media", "L"),
            "master_shapes": walk(slide.slide_layout.slide_master.shapes, sw, sh, case / "media", "M"),
            "background": {
                "slide": background_image(slide, case / "media", "bg-slide") or background_color(slide),
                "layout": background_image(slide.slide_layout, case / "media", "bg-layout") or background_color(slide.slide_layout),
                "master": background_image(slide.slide_layout.slide_master, case / "media", "bg-master")
                or background_color(slide.slide_layout.slide_master),
            },
        }
        save(case / "extract.json", extract)
        pdf_page_png(pdf, n, case / "original.png")
        write_case_config(case, lib, lib / "layouts")
        (lib / "playlists" / "case.toml").write_text(f'name = "case"\ntitle = "{name}"\nslides = ["slide"]\n')
        save(case / "case.json", {
            "name": name, "deck": str(deck), "slide": n, "renderer": "libreoffice",
            "sldr": {"library": str(lib), "playlist": "case", "config": str(case / "config")},
            "created": dt.datetime.now().isoformat(timespec="seconds"),
        })
        print(f"  + {name}  ({len(extract['shapes'])} shapes, layout '{extract['layout']}')")


def cmd_baseline(a) -> None:
    """Point a case at an existing library slide (a hand port) instead of its own lib."""
    case = case_dir(a.case)
    meta = load(case / "case.json")
    library = Path(a.library).expanduser().resolve()
    playlist = f"parity-{meta['name']}"
    pl_dir = case / "baseline-playlists"
    pl_dir.mkdir(exist_ok=True)
    flavor = f'flavor = "{a.flavor}"\n' if a.flavor else ""
    (pl_dir / f"{playlist}.toml").write_text(f'name = "{playlist}"\ntitle = "{meta["name"]}"\n{flavor}slides = ["{a.slide}"]\n')
    cfg = case / "baseline-config" / "sldr"
    cfg.mkdir(parents=True, exist_ok=True)
    seeds = Path("~/.config/sldr/flavors").expanduser()
    (cfg / "config.toml").write_text(
        "[config]\n"
        f'library = "{library}"\n'
        f'flavor_dir = "{seeds}"\n'
        f'layout_dir = "{library / "layouts"}"\n'
        'default_flavor = "default"\n\n'
        "[presentations]\n"
        f'slide_dir = "{library / "slides"}"\n'
        f'playlist_dir = "{pl_dir}"\n'
        f'output_dir = "{case / "out-baseline"}"\n'
    )
    meta["baseline"] = {"library": str(library), "slide": a.slide, "flavor": a.flavor,
                        "playlist": playlist, "config": str(case / "baseline-config")}
    save(case / "case.json", meta)
    score(case, which="baseline", note=a.note or f"hand port {a.slide}")


# ---------------------------------------------------------------- render & score


def render(case: Path, which: str = "case") -> Path:
    meta = load(case / "case.json")
    target = meta["baseline"] if which == "baseline" else meta["sldr"]
    env = dict(os.environ, XDG_CONFIG_HOME=target["config"])
    out_png = case / ("render-baseline.png" if which == "baseline" else "render.png")
    r = run(["sldr", "export", target["playlist"], "--format", "pdf",
             "-o", str(case / f"{which}.pdf")], env=env)
    pdf = case / f"{which}.pdf"
    if r.returncode != 0 or not pdf.exists():
        die("sldr export failed:\n" + (r.stdout + r.stderr).strip()[-1500:])
    pdf_page_png(pdf, 1, out_png)
    warnings = [ln.strip() for ln in (r.stdout + r.stderr).splitlines() if ln.strip().startswith(("!", "warning"))]
    if warnings:
        (case / f"{which}.warnings.txt").write_text("\n".join(warnings) + "\n")
    return out_png


def words(text: str) -> list[str]:
    return re.findall(r"[\w€$%]+", text.lower())


def extract_text(shapes: list[dict]) -> str:
    parts = []
    for s in shapes:
        for p in s.get("text", []):
            parts.append("".join(r["text"] for r in p["runs"]))
        for row in s.get("table", []) or []:
            parts.extend(row)
        parts.append(extract_text(s.get("children", [])))
    return " ".join(parts)


def sldr_text(case: Path, which: str) -> str:
    """The text the sldr slide actually renders, read back from its PDF."""
    r = run(["pdftotext", "-f", "1", "-l", "1", "-layout", str(case / f"{which}.pdf"), "-"])
    return r.stdout if r.returncode == 0 else ""


def diff_regions(err: np.ndarray, cells: int = 12, top: int = 6) -> list[dict]:
    """The worst grid cells, as % boxes, so the agent knows where to look."""
    h, w = err.shape
    out = []
    for gy in range(cells):
        for gx in range(cells):
            block = err[gy * h // cells:(gy + 1) * h // cells, gx * w // cells:(gx + 1) * w // cells]
            out.append({"x": round(gx * 100 / cells, 1), "y": round(gy * 100 / cells, 1),
                        "w": round(100 / cells, 1), "h": round(100 / cells, 1),
                        "error": round(float(block.mean()), 3)})
    out.sort(key=lambda c: -c["error"])
    return [c for c in out[:top] if c["error"] > 0.05]


def label(img: Image.Image, text: str) -> Image.Image:
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans-Bold.ttf", 28)
    except OSError:
        font = ImageFont.load_default()
    d.rectangle([0, 0, 24 + 16 * len(text), 44], fill=(0, 0, 0))
    d.text((12, 8), text, fill=(255, 255, 0), font=font)
    return img


def score(case: Path, which: str = "case", note: str | None = None) -> dict:
    from skimage.metrics import structural_similarity

    render_png = render(case, which)
    orig = Image.open(case / "original.png").convert("RGB").resize((W, H))
    mine = Image.open(render_png).convert("RGB").resize((W, H))
    a = np.asarray(orig.resize((W // 2, H // 2)).convert("L"), dtype=np.float64) / 255
    b = np.asarray(mine.resize((W // 2, H // 2)).convert("L"), dtype=np.float64) / 255
    ssim = structural_similarity(a, b, data_range=1.0)
    # Where to look: blurred brightness difference, not the SSIM map (SSIM is
    # unstable on large flat areas and lights up a whole dark background).
    # Tolerate small offsets: a pixel counts as different only if no shift of
    # up to 3 px (6 px at full size) brings the two images together.
    from skimage.filters import gaussian
    k = 3
    pad = np.pad(b, k, mode="edge")
    diff = np.full_like(a, np.inf)
    for dy in range(-k, k + 1):
        for dx in range(-k, k + 1):
            diff = np.minimum(diff, np.abs(a - pad[k + dy:k + dy + a.shape[0], k + dx:k + dx + a.shape[1]]))
    err = np.clip(gaussian(diff, sigma=2) / 0.2, 0, 1)
    near = round(float(1 - (err > 0.25).mean()), 4)
    color_err = np.abs(np.asarray(orig, dtype=np.float64) - np.asarray(mine, dtype=np.float64)).mean() / 255

    # Red only where the slides differ, in proportion to how much: the
    # original stays readable underneath.
    alpha = np.asarray(Image.fromarray((err * 255).astype(np.uint8)).resize((W, H)), dtype=np.float64)[..., None] / 255
    base = np.asarray(orig, dtype=np.float64) * 0.6
    red = np.zeros_like(base); red[..., 0] = 255
    overlay = Image.fromarray((base * (1 - alpha) + red * alpha).astype(np.uint8))
    suffix = "-baseline" if which == "baseline" else ""
    overlay.save(case / f"heatmap{suffix}.png")
    compare = Image.new("RGB", (W * 3 // 2, H // 2))
    for i, (img, txt) in enumerate(((orig, "original"), (mine, "sldr"), (overlay, "difference"))):
        compare.paste(label(img.resize((W // 2, H // 2)), txt), (i * W // 2, 0))
    compare.save(case / f"compare{suffix}.png")

    extract = load(case / "extract.json")
    want = words(extract_text(extract["shapes"]))
    have = set(words(sldr_text(case, which)))
    missing = [w for w in dict.fromkeys(want) if w not in have]
    recall = round(1 - len(missing) / len(set(want)), 3) if want else 1.0

    result = {
        "which": which, "at": dt.datetime.now().isoformat(timespec="seconds"),
        "ssim": round(float(ssim), 4), "match": near, "color_error": round(float(color_err), 4),
        "text_recall": recall, "missing_words": missing[:40],
        "worst_regions": diff_regions(err),
        "compare": str(case / f"compare{suffix}.png"), "heatmap": str(case / f"heatmap{suffix}.png"),
        "note": note,
    }
    save(case / f"score{suffix}.json", result)
    with open(case / "history.jsonl", "a") as f:
        f.write(json.dumps(result, ensure_ascii=False) + "\n")
    return result


def cmd_render(a) -> None:
    print(render(case_dir(a.case)))


def cmd_score(a) -> None:
    r = score(case_dir(a.case), note=a.note)
    print(json.dumps(r, indent=2, ensure_ascii=False))


def cmd_gap(a) -> None:
    case = case_dir(a.case)
    if a.kind not in GAP_KINDS:
        die(f"--kind must be one of {', '.join(GAP_KINDS)}")
    rec = {"at": dt.datetime.now().isoformat(timespec="seconds"), "kind": a.kind,
           "feature": a.feature, "evidence": a.evidence}
    with open(case / "gaps.jsonl", "a") as f:
        f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    print(json.dumps(rec, ensure_ascii=False))


# ---------------------------------------------------------------- report


def cmd_report(a) -> None:
    cases = sorted(p for p in (LAB / "cases").glob("*") if (p / "case.json").exists())
    rows, gaps = [], {}
    for c in cases:
        meta = load(c / "case.json")
        cells = []
        for suffix, title in (("", "agent"), ("-baseline", "hand port")):
            s = c / f"score{suffix}.json"
            if s.exists():
                r = load(s)
                cells.append(f"<div><b>{title}</b> SSIM {r['ssim']:.3f} | text {r['text_recall']:.0%}"
                             f"<br><img src='{os.path.relpath(r['compare'], LAB)}'></div>")
        g = c / "gaps.jsonl"
        if g.exists():
            for line in g.read_text().splitlines():
                rec = json.loads(line)
                key = (rec["kind"], rec["feature"])
                gaps.setdefault(key, []).append(meta["name"])
        rows.append(f"<section><h2>{html.escape(meta['name'])}</h2>{''.join(cells) or '<i>not scored</i>'}</section>")
    gap_rows = "".join(
        f"<tr><td>{html.escape(k[0])}</td><td>{html.escape(k[1])}</td><td>{len(v)}</td><td>{html.escape(', '.join(sorted(set(v))))}</td></tr>"
        for k, v in sorted(gaps.items(), key=lambda kv: -len(kv[1])))
    doc = f"""<!doctype html><meta charset=utf-8><title>sldr parity</title>
<style>body{{font:14px system-ui;margin:24px;background:#111;color:#eee}} img{{width:100%;max-width:1440px}}
table{{border-collapse:collapse}} td,th{{border:1px solid #444;padding:4px 8px}} section{{margin:24px 0}}</style>
<h1>sldr parity</h1><p>{len(cases)} cases</p>
<h2>Gaps across cases</h2><table><tr><th>kind</th><th>feature</th><th>cases</th><th>where</th></tr>{gap_rows}</table>
{''.join(rows)}"""
    out = LAB / "report.html"
    out.write_text(doc)
    print(out)


# ---------------------------------------------------------------- main


def main() -> None:
    ap = argparse.ArgumentParser(prog="parity", description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("ingest"); p.add_argument("deck"); p.add_argument("--slides", default="all")
    p.add_argument("--case-prefix"); p.add_argument("--force", action="store_true"); p.set_defaults(fn=cmd_ingest)
    p = sub.add_parser("baseline"); p.add_argument("case"); p.add_argument("--library", required=True)
    p.add_argument("--slide", required=True); p.add_argument("--flavor"); p.add_argument("--note"); p.set_defaults(fn=cmd_baseline)
    p = sub.add_parser("render"); p.add_argument("case"); p.set_defaults(fn=cmd_render)
    p = sub.add_parser("score"); p.add_argument("case"); p.add_argument("--note"); p.set_defaults(fn=cmd_score)
    p = sub.add_parser("gap"); p.add_argument("case"); p.add_argument("--kind", required=True)
    p.add_argument("--feature", required=True); p.add_argument("--evidence", required=True); p.set_defaults(fn=cmd_gap)
    p = sub.add_parser("report"); p.set_defaults(fn=cmd_report)
    a = ap.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
